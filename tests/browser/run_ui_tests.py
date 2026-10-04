"""Arnés de pruebas de interfaz con Playwright (Chromium) y un Firebase simulado.

Uso (desde la raíz del repositorio):
    pip install playwright openpyxl pillow && python -m playwright install chromium
    python tests/browser/run_ui_tests.py

El navegador se ejecuta deliberadamente en la zona horaria Asia/Tokyo para
comprobar que los cálculos usan America/Bogota y no la zona del equipo.
"""
import json
import functools
import http.server
import os
import sys
import threading
from pathlib import Path

from playwright.sync_api import sync_playwright, expect

sys.path.insert(0, str(Path(__file__).parent))
import fixture  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(os.environ.get("UI_TEST_OUT", ROOT / "tests" / "browser" / "out"))
OUT.mkdir(parents=True, exist_ok=True)
MOCK_JS = (Path(__file__).parent / "mock-firebase.js").read_text(encoding="utf-8")
FIXED_NOW = "2026-10-03T15:00:00-05:00"  # sábado 3 de octubre de 2026, 3:00 p. m. en Bogotá
PORT = 8765

RESULTS = []


def check(name, condition, detail=""):
    RESULTS.append((name, bool(condition), detail))
    print(("PASS " if condition else "FAIL ") + name + (f" — {detail}" if detail else ""))


def serve():
    class QuietHandler(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *args):
            pass

    handler = functools.partial(QuietHandler, directory=str(ROOT))
    server = http.server.ThreadingHTTPServer(("127.0.0.1", PORT), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server


def init_script(db, uid, extra=""):
    user = {"uid": uid, "email": uid + "@ieruu.test"} if uid else None
    return (
        "window.__MOCK = " + json.dumps({"db": db, "user": user, "loginAs": user, "latencyMs": 30}, ensure_ascii=False) + ";\n"
        + extra
    )


def new_page(browser, db, uid, viewport=None, extra=""):
    context = browser.new_context(
        viewport=viewport or {"width": 1280, "height": 900},
        timezone_id="Asia/Tokyo",
        locale="es-CO",
        accept_downloads=True,
        has_touch=bool(viewport and viewport["width"] < 600),
        is_mobile=bool(viewport and viewport["width"] < 600),
    )
    context.route("https://www.gstatic.com/firebasejs/**", lambda route: route.fulfill(
        status=200, content_type="text/javascript", body=MOCK_JS))
    page = context.new_page()
    page.clock.set_fixed_time(FIXED_NOW)
    page.add_init_script(init_script(db, uid, extra))
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.on("console", lambda message: message.type == "error" and errors.append(message.text))
    page.goto(f"http://127.0.0.1:{PORT}/index.html")
    page.wait_for_selector("#appView:not([hidden])", timeout=15000)
    page.wait_for_function("document.querySelector('#syncStatus').textContent.includes('en vivo')", timeout=15000)
    return context, page, errors


def open_stats(page):
    page.click("#statsButton")
    expect(page.locator("#statsDialog")).to_be_visible()


def pick_period(page, value):
    page.locator(f"#statsDialog label.chip:has(input[value='{value}'])").click()


def stat(page, selector):
    return int(page.locator(selector).inner_text())


def test_stats(browser, db):
    for role_uid in ("uid-docente", "uid-directivo", "uid-admin"):
        context, page, errors = new_page(browser, db, role_uid)
        check(f"[{role_uid}] botón Estadísticas visible", page.locator("#statsButton").is_visible())
        open_stats(page)

        pick_period(page, "hoy")
        check(f"[{role_uid}] Hoy = sábado 3 oct", "sábado, 3 de octubre de 2026" in page.inner_text("#statsRangeLabel").lower(),
              page.inner_text("#statsRangeLabel"))
        check(f"[{role_uid}] Hoy cuenta", stat(page, "#statsTotal") == fixture.expected_count(db, "2026-10-03", "2026-10-03"))

        pick_period(page, "semana")
        expected_week = fixture.expected_count(db, "2026-09-28", "2026-10-04")
        check(f"[{role_uid}] Esta semana (lun 28 sep – dom 4 oct)", stat(page, "#statsTotal") == expected_week,
              f"vista={stat(page, '#statsTotal')} esperado={expected_week}")

        pick_period(page, "mes")
        check(f"[{role_uid}] Este mes", stat(page, "#statsTotal") == fixture.expected_count(db, "2026-10-01", "2026-10-31"))

        pick_period(page, "mes-anterior")
        expected_sep = fixture.expected_count(db, "2026-09-01", "2026-09-30")
        check(f"[{role_uid}] Mes anterior (septiembre)", stat(page, "#statsTotal") == expected_sep)

        page.select_option("#statsGroupInput", "10.2")
        page.select_option("#statsExcuseInput", "validada")
        expected = fixture.expected_count(db, "2026-09-01", "2026-09-30", group="10.2", excuse="validada")
        check(f"[{role_uid}] filtros combinados grupo+excusa", stat(page, "#statsTotal") == expected,
              f"vista={stat(page, '#statsTotal')} esperado={expected}")
        page.click("#statsResetButton")
        check(f"[{role_uid}] limpiar filtros vuelve a Hoy", page.locator("input[value='hoy']").is_checked())
        if errors:
            check(f"[{role_uid}] sin errores de consola", False, "; ".join(errors[:3]))
        context.close()

    context, page, errors = new_page(browser, db, "uid-docente")
    open_stats(page)
    pick_period(page, "mes-de")
    page.select_option("#statsMonthInput", "2026-07")
    check("mes sin datos muestra 0 y mensaje", stat(page, "#statsTotal") == 0
          and "No hay llegadas tarde" in page.inner_text("#statsList"))

    pick_period(page, "semana-de")
    page.fill("#statsWeekInput", "2027-01-01")
    page.dispatch_event("#statsWeekInput", "change")
    label = page.inner_text("#statsRangeLabel")
    check("semana que cruza el año", "28 de diciembre de 2026" in label and "3 de enero de 2027" in label, label)
    check("estudiante retirado se cuenta en la semana del cambio de año", stat(page, "#statsTotal") == 2)
    check("estudiante retirado se marca", "Ya no figura en estudiantes" in page.inner_text("#statsList"))

    pick_period(page, "dia")
    page.fill("#statsDayInput", "2026-10-01")
    page.dispatch_event("#statsDayInput", "change")
    page.fill("#statsNameInput", "nusta ibanez")
    check("búsqueda sin tildes encuentra Ñusta Ibáñez", stat(page, "#statsTotal") == 1
          and "Ñusta Ibáñez Úsuga" in page.inner_text("#statsList"))
    check("franja desconocida aparece en el filtro", page.locator("#statsFranjaInput option[value='otra-franja']").count() == 1)
    page.fill("#statsNameInput", "")

    pick_period(page, "dia")
    page.fill("#statsDayInput", "2026-10-02")
    page.dispatch_event("#statsDayInput", "change")
    page.fill("#statsNameInput", "Legado")
    check("grupo heredado 102 se muestra como 10.2", "Grupo 10.2" in page.inner_text("#statsList"))
    page.fill("#statsNameInput", "")

    # Desglose: la suma por grupo coincide con el total.
    pick_period(page, "mes-anterior")
    page.check("#statsShowEmptyGroups")
    rows = page.locator("#statsGroupTableBody tr").count()
    check("con 'grupos en cero' se listan 18 grupos (+ Sin grupo si aplica)", rows >= 18, str(rows))
    total_from_rows = page.evaluate(
        "[...document.querySelectorAll('#statsGroupTableBody tr')].reduce((s, r) => s + Number(r.children[1].textContent || 0), 0)")
    check("suma del desglose = total", total_from_rows == stat(page, "#statsTotal"))
    check("paginación del listado (200 primeros)", page.locator("#statsList li").count() == 200
          and page.locator("#statsMoreButton").is_visible())
    page.click("#statsMoreButton")
    check("mostrar más amplía a 400", page.locator("#statsList li").count() == min(400, stat(page, "#statsTotal")))
    page.screenshot(path=str(OUT / "stats-desktop.png"), full_page=False)

    # Actualización en vivo: un registro nuevo hoy aparece sin recargar.
    pick_period(page, "hoy")
    before = stat(page, "#statsTotal")
    page.evaluate("""() => {
      window.__mockApi.writePath('registros/-Nlive1', {uid: 'EDGE01', nombre: 'Ñusta Ibáñez Úsuga', grupo: '6.1',
        fecha: '2026-10-03', hora: '12:15:00', timestamp: 1, periodo: 'descanso2', tarde: true});
      window.__mockApi.notify();
    }""")
    page.wait_for_timeout(300)
    check("actualización en vivo dentro de la vista", stat(page, "#statsTotal") == before + 1)
    if errors:
        check("sin errores de consola (docente)", False, "; ".join(errors[:3]))
    context.close()

    # Móvil
    context, page, errors = new_page(browser, db, "uid-docente", viewport={"width": 390, "height": 844})
    check("móvil: botón visible sin desplazamiento horizontal",
          page.locator("#statsButton").is_visible() and page.evaluate("document.documentElement.scrollWidth <= window.innerWidth"))
    open_stats(page)
    pick_period(page, "semana")
    box = page.locator("#statsDialog").bounding_box()
    check("móvil: la consulta ocupa la pantalla", box and box["width"] >= 385, str(box))
    overflow = page.evaluate("document.querySelector('#statsDialog').scrollWidth <= document.querySelector('#statsDialog').clientWidth + 1")
    check("móvil: sin desbordamiento horizontal en el diálogo", overflow)
    page.screenshot(path=str(OUT / "stats-mobile.png"), full_page=False)
    context.close()


def download_xlsx(page, name):
    with page.expect_download(timeout=20000) as info:
        page.click("#statsExportButton")
    download = info.value
    target = OUT / name
    download.save_as(str(target))
    return download.suggested_filename, target


def inspect_xlsx(path):
    import openpyxl
    workbook = openpyxl.load_workbook(path)
    summary, detail = workbook["Resumen"], workbook["Detalle"]
    total_row = [row for row in summary.iter_rows(values_only=True) if row and row[0] == "Total"][0]
    detail_rows = [row for row in detail.iter_rows(min_row=2, values_only=True) if row and row[0] is not None]
    return workbook, summary, detail, total_row, detail_rows


def libreoffice_opens(path):
    import subprocess
    import tempfile
    with tempfile.TemporaryDirectory() as folder:
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "csv", "--outdir", folder, str(path)],
            capture_output=True, text=True, timeout=120)
        csv_files = list(Path(folder).glob("*.csv"))
        return result.returncode == 0 and bool(csv_files), (csv_files[0].read_text(encoding="utf-8", errors="replace")[:300] if csv_files else result.stderr)


def test_export(browser, db):
    import datetime
    for role_uid in ("uid-docente", "uid-directivo", "uid-admin"):
        context, page, errors = new_page(browser, db, role_uid)
        open_stats(page)
        pick_period(page, "semana")
        view_total = stat(page, "#statsTotal")
        name, path = download_xlsx(page, f"semana-{role_uid}.xlsx")
        check(f"[{role_uid}] nombre del archivo", name == "llegadas-tarde_semana_2026-09-28_a_2026-10-04_generado-2026-10-03.xlsx", name)
        workbook, summary, detail, total_row, rows = inspect_xlsx(path)
        check(f"[{role_uid}] hojas Resumen y Detalle", workbook.sheetnames == ["Resumen", "Detalle"], str(workbook.sheetnames))
        check(f"[{role_uid}] total del Excel = vista", total_row[1] == view_total == len(rows), f"excel={total_row[1]} filas={len(rows)} vista={view_total}")
        if errors:
            check(f"[{role_uid}] sin errores de consola en exportación", False, "; ".join(errors[:3]))
        context.close()

    context, page, errors = new_page(browser, db, "uid-docente")
    open_stats(page)
    pick_period(page, "mes-anterior")
    page.select_option("#statsGroupInput", "6.1")
    view = {key: stat(page, sel) for key, sel in (("total", "#statsTotal"), ("sin", "#statsWithoutExcuse"), ("val", "#statsValidated"), ("est", "#statsStudents"))}
    name, path = download_xlsx(page, "septiembre-6-1.xlsx")
    workbook, summary, detail, total_row, rows = inspect_xlsx(path)
    check("mes anterior + grupo: totales idénticos (total, sin excusa, validadas, estudiantes)",
          (total_row[1], total_row[2], total_row[4], total_row[6]) == (view["total"], view["sin"], view["val"], view["est"]),
          f"excel={total_row} vista={view}")
    check("con filtro de grupo sólo aparece ese grupo en el resumen",
          [row[0] for row in summary.iter_rows(min_row=9, values_only=True) if row and row[0]] == ["Grupo 6.1", "Total"])
    first = rows[0]
    first_view = page.locator("#statsList li").first.inner_text()
    check("fecha es un valor de fecha real", isinstance(first[0], datetime.datetime), repr(first[0]))
    check("hora es un valor de hora real", isinstance(first[1], datetime.time), repr(first[1]))
    check("formato de fecha dd/mm/yyyy", detail["A2"].number_format == "dd/mm/yyyy", detail["A2"].number_format)
    check("primera fila del Excel = primera de la vista",
          first[2] in first_view and first[1].strftime("%H:%M:%S") in first_view, f"{first[:3]} vs {first_view!r}")
    check("grupo como texto, no número", detail["D2"].value == "6.1" and detail["D2"].data_type == "s", repr(detail["D2"].value))
    check("encabezado en negrita", detail["A1"].font.b is True)
    check("primera fila fija (freeze)", detail.freeze_panes == "A2", str(detail.freeze_panes))
    check("ancho de columna razonable", (detail.column_dimensions["C"].width or 0) >= 30)
    check("periodo con tildes en el resumen", summary["B2"].value == "Septiembre de 2026", repr(summary["B2"].value))
    check("rango Desde/Hasta como fechas", summary["B3"].value == datetime.datetime(2026, 9, 1) and summary["B4"].value == datetime.datetime(2026, 9, 30))
    ok, preview = libreoffice_opens(path)
    check("LibreOffice abre el archivo sin error", ok, preview[:120])

    # Nombre con ñ y tildes, día específico.
    pick_period(page, "dia")
    page.fill("#statsDayInput", "2026-10-03")
    page.dispatch_event("#statsDayInput", "change")
    page.select_option("#statsGroupInput", "")
    page.fill("#statsNameInput", "ñusta")
    name, path = download_xlsx(page, "dia-nusta.xlsx")
    workbook, summary, detail, total_row, rows = inspect_xlsx(path)
    check("tildes y ñ intactas", rows and rows[0][2] == "Ñusta Ibáñez Úsuga" and rows[0][6] == "Cita odontológica — acudiente", repr(rows[:1]))
    check("filtro de nombre descrito en el resumen", "ñusta" in summary["B5"].value, summary["B5"].value)

    # Periodo vacío.
    page.fill("#statsNameInput", "")
    pick_period(page, "mes-de")
    page.select_option("#statsMonthInput", "2026-07")
    name, path = download_xlsx(page, "vacio.xlsx")
    workbook, summary, detail, total_row, rows = inspect_xlsx(path)
    check("mes sin datos: archivo válido con total 0", total_row[1] == 0 and name.endswith("mes_2026-07_generado-2026-10-03.xlsx"), name)
    ok, _ = libreoffice_opens(path)
    check("LibreOffice abre el archivo vacío", ok)
    context.close()

    # Fallo de red al cargar la biblioteca y reintento.
    context, page, errors = new_page(browser, db, "uid-docente")
    attempts = {"count": 0}

    def flaky(route):
        attempts["count"] += 1
        if attempts["count"] == 1:
            route.abort()
        else:
            route.continue_()
    page.route("**/vendor/write-excel-file-4.1.1.min.js", flaky)
    open_stats(page)
    page.click("#statsExportButton")
    page.wait_for_function("document.querySelector('#statsExportStatus').classList.contains('error')", timeout=10000)
    check("sin conexión: mensaje de error visible", "No se pudo generar" in page.inner_text("#statsExportStatus"))
    name, path = download_xlsx(page, "reintento.xlsx")
    check("reintento tras fallo de red descarga el archivo", path.exists() and attempts["count"] == 2)
    context.close()

    # Rendimiento con un año escolar grande.
    big = fixture.build(students_count=1200, late_per_day=120, seed=3)
    context, page, errors = new_page(browser, big, "uid-docente")
    open_stats(page)
    started = page.evaluate("performance.now()")
    pick_period(page, "mes-anterior")
    elapsed = page.evaluate(f"performance.now() - {started}")
    name, path = download_xlsx(page, "grande.xlsx")
    workbook, summary, detail, total_row, rows = inspect_xlsx(path)
    check("volumen alto: consulta y exportación correctas", total_row[1] == stat(page, "#statsTotal") == len(rows),
          f"{len(big['registros'])} registros en la base, {len(rows)} en el mes, consulta {elapsed:.0f} ms")
    context.close()


def main():
    selected = set(sys.argv[1:]) or {"stats", "export", "evidence"}
    db = fixture.build()
    server = serve()
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        if "stats" in selected:
            test_stats(browser, db)
        if "export" in selected and "test_export" in globals():
            globals()["test_export"](browser, db)
        if "evidence" in selected and "test_evidence" in globals():
            globals()["test_evidence"](browser, db)
        browser.close()
    server.shutdown()
    failed = [result for result in RESULTS if not result[1]]
    print(f"\n{len(RESULTS) - len(failed)}/{len(RESULTS)} comprobaciones superadas")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
